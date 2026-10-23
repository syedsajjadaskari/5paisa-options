"""Manual Trading Journal Importer and Audit Engine.

Parses, audits, and normalizes discretionary trader workbooks, detecting
sheet-level count discrepancies, preserving provenance metadata, and
generating structured Parquet, CSV, and Excel audit artifacts.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd

TZ_IST = ZoneInfo("Asia/Kolkata")


@dataclass(frozen=True)
class JournalTradeRecord:
    """Normalized record of a manually logged trading journal entry."""

    row_number: int
    sl_no: int
    trade_date: date
    entry_time_str: str
    entry_datetime_ist: datetime
    strategy: str
    outcome: str
    source_file: str
    source_sha256: str


@dataclass
class JournalAuditReport:
    """Detailed audit metrics and discrepancy summary for the manual journal workbook."""

    source_path: Path
    source_sha256: str
    total_detailed_trades: int
    summary_reported_trades: int
    has_discrepancy: bool
    discrepancy_details: list[str]
    outcome_breakdown_detailed: dict[str, int]
    outcome_breakdown_summary: dict[str, int]
    strategy_breakdown: dict[str, int]
    min_date: date
    max_date: date
    trades: list[JournalTradeRecord] = field(default_factory=list)


class ManualJournalImporter:
    """Audits and imports the Excel manual trading journal without altering the original file."""

    def __init__(
        self, journal_path: Path | str = "data/manual/Trading_Journal_Sep2025_Aug2026 REPORT.xlsx"
    ) -> None:
        self.journal_path = Path(journal_path)

    def compute_sha256(self) -> str:
        """Calculate SHA-256 hash of the target workbook."""
        with open(self.journal_path, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()

    def audit_and_import(self) -> JournalAuditReport:
        """Parse workbook, audit sheets for discrepancies, and extract normalized records."""
        if not self.journal_path.exists():
            raise FileNotFoundError(f"Trading journal not found at {self.journal_path}")

        file_hash = self.compute_sha256()
        wb = openpyxl.load_workbook(self.journal_path, data_only=True)

        if "Trading Journal" not in wb.sheetnames or "Summary" not in wb.sheetnames:
            raise ValueError(f"Workbook missing expected sheets. Found: {wb.sheetnames}")

        # 1. Audit Summary Sheet
        ws_sum = wb["Summary"]
        summary_metrics: dict[str, Any] = {}
        for r in range(2, ws_sum.max_row + 1):
            k = ws_sum.cell(r, 1).value
            v = ws_sum.cell(r, 2).value
            if k is not None:
                summary_metrics[str(k).strip()] = v

        summary_total = int(summary_metrics.get("Total trades", 0))
        summary_outcomes = {
            "Target": int(summary_metrics.get("Target outcomes", 0)),
            "BE": int(summary_metrics.get("Break-even outcomes", 0)),
            "SL": int(summary_metrics.get("Stop-loss outcomes", 0)),
        }

        # 2. Parse Detailed Trading Journal Sheet
        ws_det = wb["Trading Journal"]
        records: list[JournalTradeRecord] = []
        detailed_outcomes: dict[str, int] = {"Target": 0, "BE": 0, "SL": 0}
        strategy_counts: dict[str, int] = {}

        for r in range(3, ws_det.max_row + 1):
            sl_no_val = ws_det.cell(r, 1).value
            date_val = ws_det.cell(r, 2).value
            time_val = ws_det.cell(r, 3).value
            strat_val = ws_det.cell(r, 4).value
            outcome_val = ws_det.cell(r, 5).value

            if sl_no_val is None or date_val is None:
                continue

            sl_no = int(sl_no_val)

            # Parse date
            if isinstance(date_val, datetime):
                trade_date = date_val.date()
            elif isinstance(date_val, date):
                trade_date = date_val
            else:
                trade_date = date.fromisoformat(str(date_val)[:10])

            # Parse entry time
            if isinstance(time_val, time):
                entry_time_str = time_val.strftime("%H:%M")
                hour, minute = time_val.hour, time_val.minute
            elif isinstance(time_val, datetime):
                entry_time_str = time_val.strftime("%H:%M")
                hour, minute = time_val.hour, time_val.minute
            else:
                time_str_clean = str(time_val).strip()
                parts = time_str_clean.split(":")
                hour, minute = int(parts[0]), int(parts[1])
                entry_time_str = f"{hour:02d}:{minute:02d}"

            entry_datetime_ist = datetime(
                trade_date.year, trade_date.month, trade_date.day, hour, minute, 0, tzinfo=TZ_IST
            )

            strategy = str(strat_val).strip().upper()
            strategy_counts[strategy] = strategy_counts.get(strategy, 0) + 1

            outcome = str(outcome_val).strip()
            if outcome.lower() in ("target", "tgt"):
                outcome_norm = "Target"
            elif outcome.lower() in ("sl", "stoploss", "stop_loss"):
                outcome_norm = "SL"
            elif outcome.lower() in ("be", "breakeven", "break_even"):
                outcome_norm = "BE"
            else:
                outcome_norm = outcome

            detailed_outcomes[outcome_norm] = detailed_outcomes.get(outcome_norm, 0) + 1

            rec = JournalTradeRecord(
                row_number=r,
                sl_no=sl_no,
                trade_date=trade_date,
                entry_time_str=entry_time_str,
                entry_datetime_ist=entry_datetime_ist,
                strategy=strategy,
                outcome=outcome_norm,
                source_file=self.journal_path.name,
                source_sha256=file_hash,
            )
            records.append(rec)

        # 3. Discrepancy Detection
        discrepancies: list[str] = []
        has_disc = False

        total_detailed = len(records)
        if total_detailed != summary_total:
            has_disc = True
            diff = total_detailed - summary_total
            discrepancies.append(
                f"Trade count mismatch: Detailed sheet contains {total_detailed} rows, "
                f"while Summary sheet claims {summary_total} trades (Delta: {diff:+d})."
            )

        for k, sum_cnt in summary_outcomes.items():
            det_cnt = detailed_outcomes.get(k, 0)
            if det_cnt != sum_cnt:
                has_disc = True
                discrepancies.append(
                    f"Outcome '{k}' mismatch: Detailed sheet has {det_cnt}, "
                    f"Summary sheet claims {sum_cnt} (Delta: {det_cnt - sum_cnt:+d})."
                )

        if total_detailed > summary_total:
            extra_trades = records[summary_total:]
            for ext in extra_trades:
                discrepancies.append(
                    f"Extra Trade #{ext.sl_no} on {ext.trade_date} at {ext.entry_time_str} "
                    f"({ext.strategy}, {ext.outcome}) at Excel row {ext.row_number} is omitted from Summary sheet."
                )

        dates = [r.trade_date for r in records]

        return JournalAuditReport(
            source_path=self.journal_path,
            source_sha256=file_hash,
            total_detailed_trades=total_detailed,
            summary_reported_trades=summary_total,
            has_discrepancy=has_disc,
            discrepancy_details=discrepancies,
            outcome_breakdown_detailed=detailed_outcomes,
            outcome_breakdown_summary=summary_outcomes,
            strategy_breakdown=strategy_counts,
            min_date=min(dates) if dates else date.min,
            max_date=max(dates) if dates else date.max,
            trades=records,
        )

    def export_artifacts(
        self, report: JournalAuditReport, output_dir: Path | str = "reports/manual_journal"
    ) -> dict[str, Path]:
        """Generate journal_audit.xlsx, manual_trades.parquet, and manual_trades.csv."""
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)

        # 1. Prepare DataFrame
        data = [
            {
                "row_number": t.row_number,
                "sl_no": t.sl_no,
                "date": t.trade_date.isoformat(),
                "entry_time": t.entry_time_str,
                "entry_timestamp_ist": t.entry_datetime_ist.isoformat(),
                "strategy": t.strategy,
                "outcome": t.outcome,
                "source_file": t.source_file,
                "source_sha256": t.source_sha256,
            }
            for t in report.trades
        ]
        df = pd.DataFrame(data)

        # 2. Export Parquet & CSV
        parquet_path = out / "manual_trades.parquet"
        csv_path = out / "manual_trades.csv"
        df.to_parquet(parquet_path, compression="zstd", index=False)
        df.to_csv(csv_path, index=False)

        # 3. Export journal_audit.xlsx
        excel_path = out / "journal_audit.xlsx"
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            # Sheet 1: Normalized Trades
            df.to_excel(writer, sheet_name="Audited Trades", index=False)

            # Sheet 2: Audit Summary & Discrepancies
            summary_rows = [
                {"Metric": "Workbook File", "Value": str(report.source_path)},
                {"Metric": "SHA-256 Checksum", "Value": report.source_sha256},
                {"Metric": "Detailed Trades Count", "Value": report.total_detailed_trades},
                {"Metric": "Summary Sheet Stated Count", "Value": report.summary_reported_trades},
                {
                    "Metric": "Discrepancy Status",
                    "Value": "DISCREPANCY DETECTED" if report.has_discrepancy else "CLEAN MATCH",
                },
                {"Metric": "Date Range", "Value": f"{report.min_date} to {report.max_date}"},
                {
                    "Metric": "Target (Detailed vs Summary)",
                    "Value": f"{report.outcome_breakdown_detailed.get('Target', 0)} vs {report.outcome_breakdown_summary.get('Target', 0)}",
                },
                {
                    "Metric": "BE (Detailed vs Summary)",
                    "Value": f"{report.outcome_breakdown_detailed.get('BE', 0)} vs {report.outcome_breakdown_summary.get('BE', 0)}",
                },
                {
                    "Metric": "SL (Detailed vs Summary)",
                    "Value": f"{report.outcome_breakdown_detailed.get('SL', 0)} vs {report.outcome_breakdown_summary.get('SL', 0)}",
                },
            ]
            for strat, cnt in sorted(report.strategy_breakdown.items()):
                summary_rows.append({"Metric": f"Strategy: {strat}", "Value": cnt})

            pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Audit Summary", index=False)

            # Sheet 3: Discrepancy Details
            disc_rows = [
                {"Index": i + 1, "Discrepancy": d} for i, d in enumerate(report.discrepancy_details)
            ]
            pd.DataFrame(disc_rows).to_excel(writer, sheet_name="Discrepancies", index=False)

        return {
            "parquet": parquet_path,
            "csv": csv_path,
            "excel": excel_path,
        }
